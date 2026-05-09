import base64
import fitz  # pymupdf for PDF
from docx import Document
from pptx import Presentation
import openpyxl
import os

# ── PDF ──────────────────────────────────────────
def extract_pdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text()
    return text[:8000]  # limit tokens

# ── Word ─────────────────────────────────────────
def extract_docx(file_bytes: bytes) -> str:
    import io
    doc = Document(io.BytesIO(file_bytes))
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return text[:8000]

# ── Excel ─────────────────────────────────────────
def extract_xlsx(file_bytes: bytes) -> str:
    import io
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    result = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        result.append(f"Sheet: {sheet}")
        for row in ws.iter_rows(values_only=True):
            row_data = [str(c) if c is not None else "" for c in row]
            if any(row_data):
                result.append(" | ".join(row_data))
    return "\n".join(result)[:8000]

# ── PowerPoint ────────────────────────────────────
def extract_pptx(file_bytes: bytes) -> str:
    import io
    prs = Presentation(io.BytesIO(file_bytes))
    text = []
    for i, slide in enumerate(prs.slides):
        text.append(f"Slide {i+1}:")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text)
    return "\n".join(text)[:8000]

# ── Image → base64 ────────────────────────────────
def encode_image(file_bytes: bytes, content_type: str) -> str:
    return base64.b64encode(file_bytes).decode('utf-8')

# ── Route file to extractor ───────────────────────
def process_file(file_bytes: bytes, filename: str, content_type: str) -> dict:
    ext = filename.lower().split('.')[-1]

    if ext == 'pdf':
        return {"type": "text", "content": extract_pdf(file_bytes), "label": "PDF"}

    elif ext == 'docx':
        return {"type": "text", "content": extract_docx(file_bytes), "label": "Word"}

    elif ext in ['xlsx', 'xls']:
        return {"type": "text", "content": extract_xlsx(file_bytes), "label": "Excel"}

    elif ext in ['pptx', 'ppt']:
        return {"type": "text", "content": extract_pptx(file_bytes), "label": "PowerPoint"}

    elif ext in ['png', 'jpg', 'jpeg', 'gif', 'webp']:
        return {
            "type": "image",
            "content": encode_image(file_bytes, content_type),
            "media_type": content_type,
            "label": "Image"
        }

    elif ext in ['mp3', 'mp4', 'wav', 'm4a', 'ogg', 'webm']:
        return {"type": "audio", "content": file_bytes, "label": "Audio/Video"}

    else:
        return {"type": "unsupported", "content": "", "label": ext}